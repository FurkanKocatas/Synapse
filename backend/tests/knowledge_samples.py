"""Small documents built in code for parser and pipeline tests."""

import io
import zipfile

import docx
import openpyxl
import pptx
from PIL import Image, ImageDraw, ImageFont


def scanned_pdf(*lines: str, dpi: int = 200) -> bytes:
    """A one-page A4 PDF that is only an image of the lines at ``dpi``, like a scan (ASCII text:
    the built-in font)."""
    page = Image.new("L", (round(8.27 * dpi), round(11.69 * dpi)), 255)
    draw = ImageDraw.Draw(page)
    font = ImageFont.load_default(size=dpi // 5)  # about 14 points
    for index, line in enumerate(lines):
        draw.text((dpi, dpi + index * dpi // 3), line, fill=0, font=font)
    buffer = io.BytesIO()
    page.save(buffer, format="PDF", resolution=dpi)
    return buffer.getvalue()


def pdf(*pages: str) -> bytes:
    """A valid PDF with the given text per page, one printed line per "\\n" (Helvetica, ASCII
    text); an empty string makes a page without text, like a scan."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        lines = " 0 -16 Td ".join(f"({line}) Tj" for line in text.split("\n"))
        content = f"BT /F1 12 Tf 72 720 Td {lines} ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.7\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def word() -> bytes:
    document = docx.Document()
    document.add_heading("Belediye Meclisi Kararı", level=1)
    document.add_paragraph("Karar No: 2026/35. Meclis üyeleri toplandı.")
    document.add_heading("Gündem", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Madde", "Karar"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "Kabul edildi"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def word_with_merged_cells() -> bytes:
    """A table whose title cell spans both columns and whose first column spans two rows."""
    document = docx.Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Performans Göstergeleri"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "P.G. 2.7.1."
    table.cell(1, 1).text, table.cell(2, 1).text = "50", "20"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def spreadsheet() -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Bütçe"
    first.append(["Kalem", "Tutar"])
    first.append(["Personel", 1250000])
    first.append([None, None])
    workbook.create_sheet("Boş")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def spreadsheet_with_title() -> bytes:
    """A price list as municipalities publish them: a title row, a note, then the table."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ücretler"
    sheet.append(["BELEDİYE ÜCRET TARİFESİ"])
    sheet.append([None, "Fiyatlara KDV dahildir."])
    sheet.append(["Hizmet", "Ücret"])
    sheet.append(["Nikah salonu", "1500"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def slides() -> bytes:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title = slide.shapes.title
    assert title is not None
    title.text = "KVKK Eğitimi"
    slide.placeholders[1].text = "Açık rıza nedir?"
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Konuşmacı notu"
    presentation.slides.add_slide(presentation.slide_layouts[5])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def zip_bomb_word() -> bytes:
    """A Word package with one member that expands about a thousandfold."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(word())) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            target.writestr(item, source.read(item))
        target.writestr("word/padding.bin", b"\x00" * (50 * 1024 * 1024))
    return buffer.getvalue()
