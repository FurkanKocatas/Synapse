"""Text extraction behind the ``Parser`` port (ADR 0010, phase 4 step 3).

``LightParser`` reads born-digital files with small libraries: PDF text layers with PDFium,
Word, Excel and PowerPoint files with their Python readers. It keeps the unit a citation needs
(page, slide, sheet) and marks pages without usable text for OCR.

Each page carries its text (what is shown and quality-checked) and its blocks (knowledge/
structure.py: what chunking reads). Word, Excel and PowerPoint files give their structure
directly; a PDF text layer is plain lines, so its blocks come from the section rules of
knowledge/headings.py.

Runs only in workers, never in the API process (ADR 0002): these files are untrusted.
"""

import re
import threading
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, Protocol

import docx
import openpyxl
import pptx
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
from docx.oxml.ns import qn
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph

from synapse.knowledge import quality
from synapse.knowledge.filetypes import MediaType
from synapse.knowledge.headings import blocks_from_text, section
from synapse.knowledge.language import language
from synapse.knowledge.structure import Block, BlockKind, Table
from synapse.knowledge.turkish import lower

PageKind = Literal["page", "slide", "sheet", "document"]

# Fewer visible characters than this on a PDF page means the text layer is missing (a scan)
# or unusable; the page goes to OCR.
MIN_TEXT_CHARS = 20
# A page mostly covered by one picture with only a few words of text layer is a scan, and its
# words are a stamp put on it (an e-signature's, a filing number), not its content: it goes to
# OCR too. On the evaluation corpus this sends 75 of 5,052 PDF pages to OCR on top of the others:
# such a scan, section covers of annual reports, and figures and tables stored as pictures, whose
# text only OCR reads. OCR costs time there; the page keeps the better text (ocr.better_text).
PICTURE_SHARE = 0.5
STAMP_WORDS = 30
MAX_PAGES = 5000
MAX_SHEET_CELLS = 500_000
# Office files are zip packages; refuse ones that expand to far more than they weigh, and ones
# with more members than real documents have (each member costs memory before it is read).
MAX_UNCOMPRESSED_BYTES = 1 << 30
MAX_COMPRESSION_RATIO = 200
MAX_PACKAGE_MEMBERS = 10_000
# A slide's title is a section of the deck; a sheet's name is the outermost section.
SLIDE_TITLE_LEVEL = 2
SHEET_LEVEL = 1
# Rows above a sheet's table with fewer filled cells are its title and notes.
MIN_HEADER_CELLS = 2
# Running headers and footers: lines this close to a page's top or bottom, repeated on at least
# half of a document's pages and on at least MIN_RUNNING_PAGES.
RUNNING_EDGE = 2
MIN_RUNNING_PAGES = 3

# PDFium is not thread-safe, and the worker parses in threads.
PDFIUM_LOCK = threading.Lock()


QualityIssue = Literal["no_text", "not_turkish_like", "ocr_artefacts"]


@dataclass(frozen=True)
class Page:
    number: int
    kind: PageKind
    text: str
    label: str | None = None
    # Why the text needs OCR; None when the text layer is usable.
    issue: QualityIssue | None = None
    char_score: float | None = None
    artefacts: float | None = None
    blocks: tuple[Block, ...] = ()

    @property
    def needs_ocr(self) -> bool:
        return self.issue is not None


def scanned_page(
    number: int,
    kind: PageKind,
    text: str,
    block_text: str | None = None,
    picture_share: float = 0.0,
) -> Page:
    """A page of a format that can carry an OCR'd text layer: PDF pages and images.

    ``block_text``: the text its blocks are made from, when it differs from what the page shows
    (running headers and footers removed). ``picture_share``: how much of the page its largest
    picture covers.
    """
    blocks = tuple(blocks_from_text(text if block_text is None else block_text, number))
    stamped = picture_share >= PICTURE_SHARE and len(text.split()) < STAMP_WORDS
    if visible_chars(text) < MIN_TEXT_CHARS or stamped:
        return Page(number, kind, text, issue="no_text", blocks=blocks)
    assessed = quality.assess(text)
    issue: QualityIssue | None = None
    if assessed.reason == "not_turkish_like":
        issue = "not_turkish_like"
    elif assessed.reason == "ocr_artefacts":
        issue = "ocr_artefacts"
    return Page(
        number,
        kind,
        text,
        issue=issue,
        char_score=assessed.char_score,
        artefacts=assessed.artefacts,
        blocks=blocks,
    )


@dataclass(frozen=True)
class Parsed:
    pages: list[Page]


class ParseError(ValueError):
    """The file cannot be read; ``reason`` is a stable code shown to editors."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


class Parser(Protocol):
    def parse(self, path: Path, media_type: MediaType) -> Parsed: ...


def normalize(text: str) -> str:
    """NFC, Unix line ends, no NUL or other control characters except newline and tab."""
    text = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C").strip()


def visible_chars(text: str) -> int:
    return sum(1 for ch in text if not ch.isspace())


class LightParser:
    def parse(self, path: Path, media_type: MediaType) -> Parsed:
        match media_type:
            case MediaType.PDF:
                return _pdf(path)
            case MediaType.DOCX:
                _check_package(path)
                return _docx(path)
            case MediaType.XLSX:
                _check_package(path)
                return _xlsx(path)
            case MediaType.PPTX:
                _check_package(path)
                return _pptx(path)
            case MediaType.PNG | MediaType.JPEG | MediaType.TIFF:
                # An image is one page with no text layer.
                return Parsed([Page(1, "page", "", issue="no_text")])


def _pdf(path: Path) -> Parsed:
    with PDFIUM_LOCK:
        try:
            document = pdfium.PdfDocument(path)
        except pdfium.PdfiumError as error:
            reason = "encrypted" if "password" in str(error).lower() else "unreadable"
            raise ParseError(reason) from error
        try:
            if len(document) > MAX_PAGES:
                raise ParseError("too_many_pages")
            texts, pictures = [], []
            for index in range(len(document)):
                page = document[index]
                textpage = page.get_textpage()
                texts.append(normalize(textpage.get_text_range()))
                pictures.append(picture_share(page))
                textpage.close()
                page.close()
        finally:
            document.close()
    if not texts:
        raise ParseError("empty")
    cleaned = without_running_lines(texts)
    return Parsed(
        [
            scanned_page(n, "page", text, block_text, share)
            for n, (text, block_text, share) in enumerate(
                zip(texts, cleaned, pictures, strict=True), start=1
            )
        ]
    )


def picture_share(page: pdfium.PdfPage) -> float:
    """The share of the page its largest picture covers (pictures inside forms included)."""
    width, height = page.get_size()
    largest = 0.0
    for picture in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE], max_depth=5):
        left, bottom, right, top = picture.get_bounds()
        largest = max(largest, max(0.0, right - left) * max(0.0, top - bottom))
    return min(1.0, largest / (width * height)) if width and height else 0.0


def without_running_lines(texts: list[str], page_word: str | None = None) -> list[str]:
    """Page texts without their page numbers and without running headers and footers after
    their first appearance (ADR 0010, ingestion rule 5).

    Only the first and last ``RUNNING_EDGE`` lines of a page are looked at. A page number
    ("3", "- 3 -", "Sayfa 3 / 40") is always dropped. Any other line found there, exactly as
    written, on at least half the pages (and at least three) is running: the publisher's name,
    the document's title. It stays where it first appears, so a signature block repeated on
    every page can still be found once. Lines differing only in their numbers are different:
    "KARAR SAYISI : 413" on one page and "KARAR SAYISI : 414" on the next are both content.
    """
    if len(texts) < MIN_RUNNING_PAGES:
        return texts
    page_number = _page_number_pattern(page_word or language().page_word)
    pages = [text.split("\n") for text in texts]
    seen_at_edges: Counter[str] = Counter()
    for lines in pages:
        seen_at_edges.update({_fold(lines[i]) for i in _edges(lines)})
    needed = max(MIN_RUNNING_PAGES, len(pages) // 2)
    running = {line for line, count in seen_at_edges.items() if count >= needed}
    kept_once: set[str] = set()
    cleaned = []
    for lines in pages:
        edges = set(_edges(lines))
        out = []
        for index, line in enumerate(lines):
            if index in edges:
                folded = _fold(line)
                if page_number.fullmatch(folded):
                    continue
                if folded in running:
                    if folded in kept_once:
                        continue
                    kept_once.add(folded)
            out.append(line)
        cleaned.append("\n".join(out))
    return cleaned


def _page_number_pattern(word: str) -> re.Pattern[str]:
    # Three digits at most, so a year alone ("2024") is not taken for a page number.
    return re.compile(rf"[-\s]*(?:{re.escape(word)}\s*)?\d{{1,3}}(?:\s*/\s*\d{{1,4}})?[-\s]*")


def _edges(lines: list[str]) -> list[int]:
    """Indices of the first and last RUNNING_EDGE lines that have text."""
    filled = [i for i, line in enumerate(lines) if line.strip()]
    return filled[:RUNNING_EDGE] + filled[RUNNING_EDGE:][-RUNNING_EDGE:]


def _fold(line: str) -> str:
    return lower(" ".join(line.split()))


def _check_package(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as package:
            entries = package.infolist()
    except zipfile.BadZipFile as error:
        raise ParseError("unreadable") from error
    if len(entries) > MAX_PACKAGE_MEMBERS:
        raise ParseError("suspicious_package")
    expanded = sum(entry.file_size for entry in entries)
    compressed = max(1, sum(entry.compress_size for entry in entries))
    if expanded > MAX_UNCOMPRESSED_BYTES or expanded / compressed > MAX_COMPRESSION_RATIO:
        raise ParseError("suspicious_package")


def _docx(path: Path) -> Parsed:
    try:
        document = docx.Document(str(path))
    except Exception as error:  # python-docx raises many kinds for a broken package
        raise ParseError("unreadable") from error
    lines: list[str] = []
    blocks: list[Block] = []
    # Paragraphs and tables in document order, headings marked so chunking can follow them.
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = " ".join(item.text.split())
            if not text:
                continue
            level = _heading_level(item)
            lines.append(f"{'#' * level} {text}" if level else text)
            blocks.append(_word_block(item, text, level))
        elif isinstance(item, WordTable):
            table = _word_table(item)
            lines.extend(_table_lines(item))
            blocks.append(Block("table", "", 1, table=table))
    return Parsed([Page(1, "document", normalize("\n".join(lines)), blocks=tuple(blocks))])


def _heading_level(paragraph: Paragraph) -> int:
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    # "Heading 2" in English templates, "Başlık 2" in Turkish ones.
    for prefix in ("Heading ", "Başlık "):
        if name.startswith(prefix) and name[len(prefix) :].isdigit():
            return min(int(name[len(prefix) :]), 6)
    return 1 if name == "Title" else 0


def _word_block(paragraph: Paragraph, text: str, style_level: int) -> Block:
    if style_level:
        # The author's heading style is the level; the document's numbering does not override it.
        return Block("heading", text, 1, level=style_level)
    found = section(text, marked=False)
    if found is not None and not found.running:
        # A heading typed as plain text ("İKİNCİ BÖLÜM" in the Normal style).
        return Block("heading", text, 1, level=found.level)
    # python-docx has no public numbering API.
    properties = paragraph._p.pPr
    numbered = properties is not None and properties.numPr is not None
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    kind: BlockKind = "list_item" if numbered or style.startswith("List") else "paragraph"
    if found is not None:
        return Block(kind, text, 1, level=found.level, label=found.label)
    return Block(kind, text, 1)


def _word_table(table: WordTable) -> Table:
    """Cell text by row; a merged cell keeps its text in its first position only."""
    rows: list[tuple[str, ...]] = []
    above: list[object] = []
    header_rows = 0
    for index, row in enumerate(table.rows):
        cells = row.cells
        texts = []
        for column, cell in enumerate(cells):
            repeats_left = column > 0 and cells[column - 1]._tc is cell._tc
            repeats_above = column < len(above) and above[column] is cell._tc
            texts.append("" if repeats_left or repeats_above else " ".join(cell.text.split()))
        above = [cell._tc for cell in cells]
        rows.append(tuple(texts))
        # Rows marked "repeat as header row" in Word.
        properties = row._tr.trPr
        marked = properties is not None and properties.find(qn("w:tblHeader")) is not None
        if index == header_rows and marked:
            header_rows += 1
    if header_rows == 0 and len(rows) > 1:
        header_rows = 1  # most Word tables label their columns in the first row, unmarked
    return Table(tuple(rows), header_rows)


def _table_lines(table: WordTable) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        # Merged cells repeat; keep each distinct value once, in order.
        deduped = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
        if deduped:
            lines.append(" | ".join(deduped))
    return lines


def _xlsx(path: Path) -> Parsed:
    # Opened through a handle: given a path, openpyxl insists on an .xlsx name, and blobs are
    # stored under their hash.
    handle = path.open("rb")
    try:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many kinds for a broken package
        handle.close()
        raise ParseError("unreadable") from error
    pages = []
    try:
        for number, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[tuple[str, ...]] = []
            cells = 0
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_SHEET_CELLS:
                    raise ParseError("sheet_too_large")
                values = ["" if value is None else str(value).strip() for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append(tuple(values))
            text = normalize("\n".join("\t".join(row) for row in rows))
            blocks = _sheet_blocks(sheet.title, number, rows)
            pages.append(Page(number, "sheet", text, label=sheet.title, blocks=blocks))
    finally:
        workbook.close()
        handle.close()
    return Parsed(pages)


def _sheet_blocks(name: str, number: int, rows: list[tuple[str, ...]]) -> tuple[Block, ...]:
    """The sheet's name as its heading, title and note rows above the table as paragraphs, and
    the table, whose header is its first row with more than one filled cell."""
    blocks = [Block("heading", name, number, level=SHEET_LEVEL)]
    lead = 0
    while lead < len(rows) and sum(1 for cell in rows[lead] if cell) < MIN_HEADER_CELLS:
        blocks.append(Block("paragraph", " ".join(c for c in rows[lead] if c), number))
        lead += 1
    body = rows[lead:]
    if body:
        table = Table(tuple(body), header_rows=1 if len(body) > 1 else 0)
        blocks.append(Block("table", "", number, table=table))
    return tuple(blocks)


def _pptx(path: Path) -> Parsed:
    try:
        presentation = pptx.Presentation(str(path))
    except Exception as error:  # python-pptx raises many kinds for a broken package
        raise ParseError("unreadable") from error
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        blocks: list[Block] = []
        title = slide.shapes.title
        for shape in slide.shapes:
            if shape.has_text_frame:
                frame = [p.strip() for p in shape.text_frame.text.split("\n") if p.strip()]
                lines.extend(frame)
                if frame and title is not None and shape.shape_id == title.shape_id:
                    text = " ".join(frame)
                    blocks.append(Block("heading", text, number, level=SLIDE_TITLE_LEVEL))
                else:
                    blocks.extend(blocks_from_text("\n".join(frame), number))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [tuple(" ".join(c.text.split()) for c in r.cells) for r in shape.table.rows]
                lines.extend(" | ".join(row) for row in rows if any(row))
                header = 1 if shape.table.first_row and len(rows) > 1 else 0
                blocks.append(Block("table", "", number, table=Table(tuple(rows), header)))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)
                blocks.extend(blocks_from_text(notes, number))
        pages.append(Page(number, "slide", normalize("\n".join(lines)), blocks=tuple(blocks)))
    return Parsed(pages)
