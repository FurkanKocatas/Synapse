"""The pilot evaluation's sheets, metrics and report.

uv run --directory backend pytest ../eval/pilot -q
"""

# ruff: noqa: S101, PLR2004, FBT001, FBT003  (pytest asserts, expected values and flags)
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from report import THRESHOLDS, metrics, render
from run import matches
from sheets import (
    CRITICAL_COLUMN,
    FOUND_COLUMN,
    FULL,
    GRADE_COLUMN,
    NO,
    PARTIAL,
    REFUSED,
    WRONG,
    YES,
    Graded,
    Question,
    Result,
    read_grading,
    read_questions,
    write_grading,
    write_question_template,
)


def _filled(tmp_path: Path, rows: list[list[object]]) -> Path:
    path = tmp_path / "questions.xlsx"
    write_question_template(path, rows=len(rows))
    book = load_workbook(path)
    sheet = book["Sorular"]
    for index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, value)
    book.save(path)
    return path


def test_the_questions_the_lawyers_write_are_read_back(tmp_path: Path) -> None:
    path = _filled(
        tmp_path,
        [
            [1, "Kira kaç TL tespit edildi?", "22.050 TL", "Gerekçeli Karar.pdf", YES, "12,5"],
            [2, "İstinaf ne karar verdi?", None, None, "hayır", None],
            [3, None, None, None, None, None],
        ],
    )
    questions = read_questions(path)
    assert [q.number for q in questions] == [1, 2]
    assert questions[0].answerable and questions[0].manual_minutes == 12.5
    assert not questions[1].answerable and questions[1].document == ""


def test_a_question_without_yes_or_no_is_refused(tmp_path: Path) -> None:
    path = _filled(tmp_path, [[1, "Soru?", "Cevap", "Belge.pdf", "belki", None]])
    with pytest.raises(ValueError, match="third column"):
        read_questions(path)


def test_two_questions_with_one_number_are_refused(tmp_path: Path) -> None:
    path = _filled(tmp_path, [[1, "Bir?", "", "", NO, None], [1, "İki?", "", "", NO, None]])
    with pytest.raises(ValueError, match="share a number"):
        read_questions(path)


def test_a_sheet_whose_headings_changed_is_refused(tmp_path: Path) -> None:
    path = _filled(tmp_path, [[1, "Soru?", "", "", NO, None]])
    book = load_workbook(path)
    book["Sorular"].cell(1, 3, "Cevap")
    book.save(path)
    with pytest.raises(ValueError, match="headings"):
        read_questions(path)


def _question(number: int, *, answerable: bool = True, minutes: float | None = None) -> Question:
    return Question(
        number, f"Soru {number}?", "cevap" if answerable else "", "Belge.pdf", answerable, minutes
    )


def test_the_grading_sheet_comes_back_with_the_lawyers_grades(tmp_path: Path) -> None:
    questions = [_question(1), _question(2, answerable=False)]
    path = tmp_path / "grading.xlsx"
    write_grading(
        path,
        [
            Result(questions[0], "answered", "Kira 22.050 TL. [1]", ["Karar, s. 3"], True, 31.2),
            Result(questions[1], "not_found", "", [], None, 6.0),
        ],
    )
    book = load_workbook(path)
    sheet = book["Puanlama"]
    assert sheet.cell(3, GRADE_COLUMN).value == REFUSED
    sheet.cell(2, GRADE_COLUMN, PARTIAL)
    sheet.cell(2, FOUND_COLUMN, YES)
    sheet.cell(2, CRITICAL_COLUMN, "Tutar ya da hesap")
    book.save(path)
    graded = read_grading(path, {q.number: q for q in questions})
    assert (
        graded[0].grade == PARTIAL and graded[0].found and graded[0].critical == "Tutar ya da hesap"
    )
    assert graded[0].seconds == 31.2
    assert graded[1].grade == REFUSED and graded[1].critical is None


def test_an_ungraded_row_is_an_error(tmp_path: Path) -> None:
    question = _question(1)
    path = tmp_path / "grading.xlsx"
    write_grading(path, [Result(question, "answered", "x", [], True, 1.0)])
    with pytest.raises(ValueError, match="grade must be"):
        read_grading(path, {1: question})


def _rows(
    answerable: tuple[int, int, int, int], unanswerable: tuple[int, int], critical: int = 0
) -> list[Graded]:
    """Graded rows: answerable ones (full, partial, wrong, refused), unanswerable ones (refused,
    answered), and the first ``critical`` rows with a critical error."""
    counts = zip((FULL, PARTIAL, WRONG, REFUSED), answerable, strict=True)
    grades = [grade for grade, count in counts for _ in range(count)]
    held, missed = unanswerable
    rows = [
        Graded(
            _question(n, minutes=10.0 if n <= 10 else None),
            "a",
            "s",
            30.0,
            g,
            True,
            "Tutar ya da hesap" if n <= critical else None,
        )
        for n, g in enumerate(grades, start=1)
    ]
    start = len(rows) + 1
    rows += [
        Graded(_question(start + n, answerable=False), "", "", 5.0, g, None, None)
        for n, g in enumerate([REFUSED] * held + [WRONG] * missed)
    ]
    return rows


def test_the_metrics_count_each_class_on_its_own() -> None:
    lines = {x.label: x for x in metrics(_rows((35, 6, 4, 5), (9, 1)), THRESHOLDS)}
    assert lines["Tam doğru cevap"].value == "%70" and lines["Tam doğru cevap"].passed
    assert lines["Tam ya da kısmen doğru cevap"].value == "%82"
    assert lines["Yanlış cevap"].value == "%8" and lines["Yanlış cevap"].passed
    assert lines["Cevabı olan sorulardan reddedilen"].value == "%10"
    assert lines["Cevabı olmayan sorulardan reddedilen"].value == "%90"
    assert lines["Bulma süresinde kazanç"].value == "%95"


def test_more_than_one_critical_error_in_fifty_fails_the_test() -> None:
    one = {x.label: x for x in metrics(_rows((45, 5, 0, 0), (10, 0), critical=1), THRESHOLDS)}
    two = {x.label: x for x in metrics(_rows((45, 5, 0, 0), (10, 0), critical=2), THRESHOLDS)}
    assert one["Kritik hukuki hata"].passed
    assert not two["Kritik hukuki hata"].passed


def test_the_report_lists_every_critical_error_with_its_question() -> None:
    rows = _rows((3, 0, 1, 0), (1, 0), critical=1)
    text = render(rows, metrics(rows, THRESHOLDS))
    assert "**Soru 1** (Tutar ya da hesap): Soru 1?" in text
    assert "| 5 | Soru 5? | Reddedildi |" in text


@pytest.mark.parametrize(
    ("expected", "title", "same"),
    [
        (
            "2024-733 Aktüerya Bilirkişi Raporu 18.03.2026.pdf",
            "2024-733 Aktüerya Bilirkişi Raporu 18.03.2026",
            True,
        ),
        ("Karar.PDF", "karar", True),
        ("Dava Dilekçesi.docx", "Cevap Dilekçesi", False),
        ("", "Karar", False),
    ],
)
def test_the_expected_document_is_found_by_its_file_name(
    expected: str, title: str, same: bool
) -> None:
    assert matches(expected, title) is same
